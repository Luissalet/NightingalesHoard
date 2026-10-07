"""Conservative OOXML workbook export that changes only imported tabular cells.

The workbook package is copied part by part. Only the selected worksheet XML and
workbook recalculation settings are edited; every other uncompressed ZIP member
is retained byte for byte. This deliberately supports a narrow, auditable slice
instead of round-tripping the whole workbook through an object model.
"""
from __future__ import annotations

import hashlib
import math
import posixpath
import re
import secrets
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, to_excel

from .workbench.engine import DataError

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"m": NS_MAIN, "r": NS_REL, "p": NS_PKG_REL}
_CELL_REF = re.compile(r"\br=['\"]([^'\"]+)['\"]")


def _sha_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _resolve_part(source_part: str, target: str) -> str:
    """Resolve a package relationship Target from its source part."""
    if target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def inventory(path: Path) -> dict[str, Any]:
    """Capture a source snapshot and useful workbook structure without saving it."""
    path = Path(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise DataError("workbook inventory requires an .xlsx or .xlsm file")
    if not zipfile.is_zipfile(path):
        raise DataError("workbook is not a valid OOXML ZIP package")
    try:
        with zipfile.ZipFile(path, "r") as package:
            names = set(package.namelist())
            workbook_xml = ET.fromstring(package.read("xl/workbook.xml"))
            properties = workbook_xml.find("m:workbookPr", NS)
            date_system = "1904" if properties is not None and properties.attrib.get("date1904", "0").lower() in {"1", "true"} else "1900"
            rels_xml = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
            targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels_xml.findall("p:Relationship", NS)}
            sheets = []
            for sheet in workbook_xml.findall("m:sheets/m:sheet", NS):
                target = targets.get(sheet.attrib.get(f"{{{NS_REL}}}id"))
                if not target:
                    continue
                part = _resolve_part("xl/workbook.xml", target)
                sheet_xml = ET.fromstring(package.read(part))
                tables = []
                comments = []
                rel_part = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
                if rel_part in names:
                    sheet_rels = ET.fromstring(package.read(rel_part))
                    related = list(sheet_rels.findall("p:Relationship", NS))
                    for relation in related:
                        rel_type = relation.attrib.get("Type", "")
                        if not (rel_type.endswith("/table") or rel_type.endswith("/comments")):
                            continue
                        resolved = _resolve_part(part, relation.attrib["Target"])
                        if rel_type.endswith("/table") and resolved in names:
                            table_root = ET.fromstring(package.read(resolved))
                            tables.append({"name": table_root.attrib.get("name"), "ref": table_root.attrib.get("ref"),
                                           "part": resolved})
                        elif rel_type.endswith("/comments") and resolved in names:
                            comments_root = ET.fromstring(package.read(resolved))
                            comments.append({"part": resolved, "refs": [item.attrib.get("ref") for item in
                                                                           comments_root.findall("m:commentList/m:comment", NS)]})
                dimension = sheet_xml.find("m:dimension", NS)
                sheets.append({"name": sheet.attrib["name"], "part": part,
                               "dimension": dimension.attrib.get("ref") if dimension is not None else None,
                               "tables": tables, "comments": comments})
            part_hashes = {name: _sha_bytes(package.read(name)) for name in sorted(names) if not name.endswith("/")}
    except (KeyError, OSError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise DataError(f"could not inventory workbook package: {exc}") from exc
    stat = path.stat()
    return {"format": path.suffix.lower(), "sha256": sha256_file(path), "bytes": stat.st_size,
            "date_system": date_system, "sheet_count": len(sheets), "sheets": sheets, "package_parts": part_hashes,
            "package_part_count": len(part_hashes)}


def _column_number(cell_ref: str) -> int:
    return coordinate_to_tuple(cell_ref)[1]


def _range_hits(ref: str, min_col: int, min_row: int, max_col: int, max_row: int) -> bool:
    try:
        a, b, c, d = range_boundaries(ref)
    except (TypeError, ValueError):
        return False
    return a <= max_col and c >= min_col and b <= max_row and d >= min_row


def _guard_target(root: ET.Element, changed_refs: set[str]) -> None:
    for cell in root.findall(".//m:sheetData/m:row/m:c", NS):
        ref = cell.attrib.get("r", "")
        if ref not in changed_refs:
            continue
        if cell.find("m:f", NS) is not None:
            raise DataError(f"preserving export refused: changed value would overwrite formula cell {ref}")
        if any(child.tag not in {f"{{{NS_MAIN}}}v", f"{{{NS_MAIN}}}is"} for child in cell):
            raise DataError(f"preserving export refused: cell {ref} contains an unsupported OOXML extension")

    for merge in root.findall("m:mergeCells/m:mergeCell", NS):
        if any(_range_hits(merge.attrib.get("ref", ""), col, row, col, row)
               for ref in changed_refs for row, col in [coordinate_to_tuple(ref)]):
            raise DataError(f"preserving export refused: changed cell intersects merged range {merge.attrib.get('ref')}")


def _excel_value(value: Any, date_system: str = "1900") -> tuple[str, str | None]:
    if value is None:
        return "", None
    if isinstance(value, bool):
        return ("1" if value else "0"), "b"
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.replace(tzinfo=None)
        epoch = CALENDAR_MAC_1904 if date_system == "1904" else CALENDAR_WINDOWS_1900
        value = to_excel(value, epoch)
    elif isinstance(value, date):
        epoch = CALENDAR_MAC_1904 if date_system == "1904" else CALENDAR_WINDOWS_1900
        value = to_excel(datetime.combine(value, time()), epoch)
    elif isinstance(value, time):
        value = (value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000) / 86400
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise DataError("preserving export cannot write non-finite Decimal values to XLSX")
        return format(value, "f"), None
    if isinstance(value, int):
        return str(value), None
    if isinstance(value, float):
        if not math.isfinite(value):
            raise DataError("preserving export cannot write non-finite floating-point values to XLSX")
        return repr(value), None
    return str(value), "inlineStr"


def _same_value(left: Any, right: Any) -> bool:
    if left is None or right is None:
        return left is None and right is None
    if isinstance(left, bool) != isinstance(right, bool):
        return False
    if isinstance(left, float) and isinstance(right, float) and math.isnan(left) and math.isnan(right):
        return True
    try:
        result = left == right
        return bool(result) if isinstance(result, bool) else False
    except (TypeError, ValueError):
        return False


def _cell_xml(opening: str, value: Any, date_system: str) -> str:
    content, kind = _excel_value(value, date_system)
    normalized_opening = re.sub(r"/\s*>$", ">", opening)
    opening_match = re.match(r"<(?P<qname>(?:[A-Za-z_][\w.-]*:)?[A-Za-z_][\w.-]*)\b[^>]*>", normalized_opening)
    if opening_match is None:
        raise DataError("preserving export refused: malformed worksheet cell tag")
    qname = opening_match.group("qname")
    attrs = re.sub(r"\s+t=['\"][^'\"]*['\"]", "", normalized_opening[:-1])
    if kind:
        attrs += f' t="{kind}"'
    if kind == "inlineStr":
        if any(not (char in "\t\n\r" or 0x20 <= ord(char) <= 0xD7FF or
                    0xE000 <= ord(char) <= 0xFFFD or 0x10000 <= ord(char) <= 0x10FFFF)
               for char in content):
            raise DataError("preserving export cannot write text containing XML-illegal control characters")
        escaped = escape(content)
        escaped = escaped.replace("\r", "&#13;")
        preserve = ' xml:space="preserve"' if content[:1].isspace() or content[-1:].isspace() else ""
        prefix = qname.rpartition(":")[0]
        inner_qname = f"{prefix}:" if prefix else ""
        inner = f"<{inner_qname}is><{inner_qname}t{preserve}>{escaped}</{inner_qname}t></{inner_qname}is>"
    elif content:
        prefix = qname.rpartition(":")[0]
        value_qname = f"{prefix}:v" if prefix else "v"
        inner = f"<{value_qname}>{content}</{value_qname}>"
    else:
        inner = ""
    return f"{attrs}>{inner}</{qname}>"


_XML_QNAME = r"(?:[A-Za-z_][\w.-]*:)?[A-Za-z_][\w.-]*"
_SHEET_DATA = re.compile(r"<(?P<qname>(?:[A-Za-z_][\w.-]*:)?sheetData)\b[^>]*>(?P<body>.*?)</(?P=qname)>", re.DOTALL)
_ROW_TAG = re.compile(r"<(?P<qname>(?:[A-Za-z_][\w.-]*:)?row)\b[^>]*(?:/>|>.*?</(?P=qname)>)", re.DOTALL)
_CELL_TAG = re.compile(r"<(?P<qname>(?:[A-Za-z_][\w.-]*:)?c)\b[^>]*(?:/>|>.*?</(?P=qname)>)", re.DOTALL)


def _cell_tag_name(cells, row_qname: str) -> str:
    for match in cells:
        cell_match = re.match(r"<(?P<qname>" + _XML_QNAME + r")\b", match.group(0))
        if cell_match:
            return cell_match.group("qname")
    row_prefix = row_qname.rpartition(":")[0]
    return f"{row_prefix}:c" if row_prefix else "c"


def _cell_ref(cell_xml: str) -> str | None:
    opening = re.match(r"<(?P<qname>" + _XML_QNAME + r")\b[^>]*>", cell_xml) or re.match(r"<(?P<qname>" + _XML_QNAME + r")\b[^>]*/>", cell_xml)
    if not opening:
        return None
    match = _CELL_REF.search(opening.group(0))
    return match.group(1) if match else None


def _patch_row(row_xml: str, replacements: dict[str, Any], date_system: str) -> str:
    opening = re.match(r"<(?P<qname>" + _XML_QNAME + r")\b[^>]*>", row_xml)
    if opening is None:
        raise DataError("preserving export refused: target worksheet row is self-closing")
    closing = re.search(r"</" + re.escape(opening.group("qname")) + r">\s*$", row_xml)
    closing_start = closing.start() if closing else -1
    if closing_start < opening.end():
        raise DataError("preserving export refused: malformed target worksheet row")
    inner = row_xml[opening.end():closing_start]
    cells = list(_CELL_TAG.finditer(inner))
    wanted = sorted(replacements.items(), key=lambda pair: _column_number(pair[0]))
    pending = iter(wanted)
    next_item = next(pending, None)
    pieces: list[str] = []
    cursor = 0
    last_cell_end = 0
    for match in cells:
        ref = _cell_ref(match.group(0))
        if ref is None:
            continue
        cell_col = _column_number(ref)
        while next_item and _column_number(next_item[0]) < cell_col:
            pieces.append(inner[cursor:match.start()])
            ref_to_add, value = next_item
            pieces.append(_cell_xml(f'<{_cell_tag_name(cells, opening.group("qname"))} r="{ref_to_add}">', value, date_system))
            cursor = match.start()
            next_item = next(pending, None)
        pieces.append(inner[cursor:match.start()])
        if ref in replacements:
            opening_match = re.match(r"<(?P<qname>" + _XML_QNAME + r")\b[^>]*>", match.group(0))
            if opening_match is None:
                opening_text = re.sub(r"/\s*>$", ">", match.group(0))
                opening_text = re.match(r"<(?P<qname>" + _XML_QNAME + r")\b[^>]*>", opening_text).group(0)
            else:
                opening_text = opening_match.group(0)
            pieces.append(_cell_xml(opening_text, replacements[ref], date_system))
            if next_item and next_item[0] == ref:
                next_item = next(pending, None)
        else:
            pieces.append(match.group(0))
        cursor = match.end()
        last_cell_end = cursor
    tail = inner[cursor:]
    if next_item:
        # Put new cells after the last cell (or at the row start) and before
        # row-level trailing content such as extension markup.
        insertion_point = last_cell_end
        tail = inner[insertion_point:]
        for ref, value in [next_item, *list(pending)]:
            pieces.append(_cell_xml(f'<{_cell_tag_name(cells, opening.group("qname"))} r="{ref}">', value, date_system))
    pieces.append(tail)
    return row_xml[:opening.end()] + "".join(pieces) + row_xml[closing_start:]


def _patch_worksheet_cells(xml: str, replacements: dict[str, Any], date_system: str) -> str:
    """Rewrite only changed cells in one pass over sheetData, preserving raw XML elsewhere."""
    if not replacements:
        return xml
    target_rows: dict[int, dict[str, Any]] = {}
    for ref, value in replacements.items():
        row, _ = coordinate_to_tuple(ref)
        target_rows.setdefault(row, {})[ref] = value
    match = _SHEET_DATA.search(xml)
    if not match:
        raise DataError("preserving export refused: worksheet has no sheetData section")
    content = match.group(2)
    output: list[str] = []
    cursor = 0
    found: set[int] = set()
    for row_match in _ROW_TAG.finditer(content):
        opening = re.match(r"<(?P<qname>" + _XML_QNAME + r")\b[^>]*>", row_match.group(0))
        row_number = int(re.search(r"\br=['\"](\d+)['\"]", opening.group(0)).group(1)) if opening else None
        if row_number in target_rows:
            output.append(content[cursor:row_match.start()])
            output.append(_patch_row(row_match.group(0), target_rows[row_number], date_system))
            found.add(row_number)
            cursor = row_match.end()
    output.append(content[cursor:])
    missing = sorted(set(target_rows) - found)
    if missing:
        raise DataError(f"preserving export refused: source row {missing[0]} is absent from worksheet XML")
    return xml[:match.start(2)] + "".join(output) + xml[match.end(2):]


def _set_full_recalc(xml: str) -> str:
    match = re.search(r"<(?P<qname>(?:[A-Za-z_][\w.-]*:)?calcPr)\b[^>]*/?>", xml)
    if match:
        tag = match.group(0)
        for key, value in (("calcMode", "auto"), ("fullCalcOnLoad", "1"), ("forceFullCalc", "1")):
            if re.search(r"\b" + key + r"=", tag):
                tag = re.sub(r"\b" + key + r"=(['\"])[^'\"]*\1", f'{key}="{value}"', tag)
            else:
                tag = tag[:-2] + f' {key}="{value}"/>' if tag.endswith("/>") else tag[:-1] + f' {key}="{value}">'
        return xml[:match.start()] + tag + xml[match.end():]
    root = re.search(r"<(?P<qname>(?:[A-Za-z_][\w.-]*:)?workbook)\b", xml)
    if not root:
        raise DataError("preserving export refused: workbook XML root is missing")
    root_qname = root.group("qname")
    closing = re.search(r"</" + re.escape(root_qname) + r">", xml)
    if not closing:
        raise DataError("preserving export refused: workbook XML closing root is missing")
    prefix = root_qname.rpartition(":")[0]
    calc_qname = f"{prefix}:calcPr" if prefix else "calcPr"
    recalc = f'<{calc_qname} calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>'
    return xml[:closing.start()] + recalc + xml[closing.start():]


def export_many_copy(source_path: Path, output_path: Path, snapshot: dict, updates: list[dict[str, Any]]) -> dict[str, Any]:
    """Preflight and patch several imported sheets into one workbook copy."""
    source_path, output_path = Path(source_path), Path(output_path)
    if source_path.suffix.lower() != ".xlsx" or snapshot.get("format") != ".xlsx" or not snapshot.get("sha256"):
        raise DataError("workbook-preserving export currently supports imported .xlsx files only; .xlsm is unsupported")
    if source_path.resolve() == output_path.resolve():
        raise DataError("workbook-preserving export cannot overwrite the source workbook")
    current_sha = sha256_file(source_path)
    if current_sha != snapshot["sha256"]:
        raise DataError("source workbook changed after ingestion; re-ingest it before preserving export")
    if not updates:
        raise DataError("workbook_updates must contain at least one dataset update")
    names = [u["sheet_name"] for u in updates]
    if len(set(names)) != len(names):
        raise DataError("workbook_updates contains duplicate sheet updates; combine changes into one dataset version")

    sheets = {item.get("name"): item for item in snapshot.get("sheets", [])}
    prepared: list[dict[str, Any]] = []
    # Validate every update before creating a temporary output or touching an existing one.
    for update in updates:
        sheet_name = update["sheet_name"]
        sheet = sheets.get(sheet_name)
        if not sheet:
            raise DataError(f"sheet {sheet_name!r} is not in the recorded source workbook inventory")
        header_row, row_count = int(update["header_row"]), int(update["row_count"])
        columns, rows, original_rows = update["columns"], update["rows"], update["original_rows"]
        if len(rows) != row_count or len(original_rows) != row_count:
            raise DataError(f"preserving export requires the imported row count ({row_count}); selected version has {len(rows)} rows")
        if any(len(r) != len(columns) for r in rows) or any(len(r) != len(columns) for r in original_rows):
            raise DataError(f"preserving export found an inconsistent row width for {update.get('dataset')}: {len(columns)} columns, current {len(rows[0]) if rows else 0}, baseline {len(original_rows[0]) if original_rows else 0}")
        if row_count > 100_000 or len(columns) > 16_384:
            raise DataError("preserving export exceeds its 100,000 row / 16,384 column safety limit")
        replacements: dict[str, Any] = {}
        for row_offset, (before, after) in enumerate(zip(original_rows, rows), start=header_row + 1):
            for col_index, (before_value, after_value) in enumerate(zip(before, after), start=1):
                if not _same_value(before_value, after_value):
                    replacements[f"{get_column_letter(col_index)}{row_offset}"] = after_value
        prepared.append({**update, "sheet": sheet, "replacements": replacements,
                         "min_row": header_row, "max_row": header_row + row_count,
                         "max_col": len(columns)})

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(output_path.name + ".tmp-" + secrets.token_hex(6))
    try:
        with zipfile.ZipFile(source_path, "r") as src:
            names_in_zip = set(src.namelist())
            patched_parts: dict[str, bytes] = {}
            for item in prepared:
                sheet = item["sheet"]
                if sheet["part"] not in names_in_zip:
                    raise DataError("source worksheet part is missing; re-ingest the workbook")
                raw = src.read(sheet["part"])
                root = ET.fromstring(raw)
                min_row, max_row, max_col = item["min_row"], item["max_row"], item["max_col"]
                _guard_target(root, set(item["replacements"]))
                for table in sheet.get("tables", []):
                    if _range_hits(table.get("ref") or "", 1, min_row, max_col, max_row):
                        if range_boundaries(table["ref"]) != (1, min_row, max_col, max_row):
                            raise DataError(f"preserving export refused: table {table.get('name')} range {table['ref']} is not exactly the imported rectangle")
                for auto_filter in root.findall("m:autoFilter", NS):
                    ref = auto_filter.attrib.get("ref", "")
                    if _range_hits(ref, 1, min_row, max_col, max_row) and range_boundaries(ref) != (1, min_row, max_col, max_row):
                        raise DataError(f"preserving export refused: worksheet filter {ref} is not exactly the imported rectangle")
                item["patched"] = _patch_worksheet_cells(raw.decode("utf-8"), item["replacements"], snapshot.get("date_system", "1900")).encode("utf-8")
                item["patched_root"] = ET.fromstring(item["patched"])
                patched_parts[sheet["part"]] = item["patched"]
            patched_parts["xl/workbook.xml"] = _set_full_recalc(src.read("xl/workbook.xml").decode("utf-8")).encode("utf-8")
            parsed_workbook_root = ET.fromstring(patched_parts["xl/workbook.xml"])
            if parsed_workbook_root is None:
                raise DataError("workbook XML validation failed")
            with zipfile.ZipFile(temp_path, "w") as dst:
                for info in src.infolist():
                    dst.writestr(info, patched_parts.get(info.filename, src.read(info.filename)))
        if sha256_file(source_path) != current_sha:
            raise DataError("source workbook changed during export; no copy was published")
        temp_path.replace(output_path)
    except (OSError, zipfile.BadZipFile, ET.ParseError, UnicodeDecodeError, KeyError) as exc:
        raise DataError(f"workbook-preserving export failed: {exc}") from exc
    finally:
        temp_path.unlink(missing_ok=True)

    with zipfile.ZipFile(output_path, "r") as check:
        output_hashes = {name: _sha_bytes(check.read(name)) for name in check.namelist() if not name.endswith("/")}
    changed_parts = sorted(name for name, digest in snapshot["package_parts"].items()
                           if name not in output_hashes or output_hashes[name] != digest)
    receipts = []
    for item in prepared:
        sheet = item["sheet"]
        root = item["patched_root"]
        receipts.append({"dataset": item["dataset"], "dataset_version": item["version"],
                         "source_id": item["source_id"], "sheet": item["sheet_name"],
                         "data_range": f"A{item['min_row'] + 1}:{get_column_letter(item['max_col'])}{item['max_row']}",
                         "header_row": item["min_row"], "row_count": item["row_count"],
                         "columns": item["columns"], "cells_written": len(item["replacements"]),
                         "preserved_metadata": {"table_count": len(sheet.get("tables", [])),
                             "comment_count": sum(len(v.get("refs", [])) for v in sheet.get("comments", [])),
                             "hyperlink_count": len(root.findall("m:hyperlinks/m:hyperlink", NS)),
                             "data_validation_count": len(root.findall("m:dataValidations/m:dataValidation", NS)),
                             "data_validation_policy": "Metadata is retained; edited values are not evaluated against validation rules."}})
    primary = receipts[0]
    return {"path": str(output_path), "format": "xlsx", "mode": "preserve_workbook",
            **primary, "cells_written": sum(item["cells_written"] for item in receipts),
            "sheets": receipts, "source_workbook": str(source_path), "source_sha256": current_sha,
            "output_sha256": sha256_file(output_path), "date_system": snapshot.get("date_system", "1900"),
            "package_part_count": len(snapshot["package_parts"]), "changed_package_parts": changed_parts,
            "unchanged_package_parts": sorted(name for name, digest in snapshot["package_parts"].items()
                                                if name in output_hashes and output_hashes[name] == digest),
            "recalculation_policy": "Workbook is marked for full recalculation on next open. Nightingale does not evaluate formulas; cached formula results may remain stale until Excel/LibreOffice recalculates.",
            "warnings": ["Only changed values are written; unchanged source cells retain their original OOXML types and content.",
                         "Comments, hyperlinks and data-validation metadata are retained; Nightingale does not enforce validation rules.",
                         "Changing a formula cell is refused; unchanged formula cells remain intact.",
                         "Workbook-level calc flags are updated to request recalculation; formula caches are not rewritten.",
                         "Macro-enabled .xlsm, row/column count changes and renamed/dropped columns are unsupported."]}


def export_copy(source_path: Path, output_path: Path, snapshot: dict, *, sheet_name: str,
                header_row: int, row_count: int, columns: list[str], rows: list[tuple[Any, ...]],
                original_rows: list[tuple[Any, ...]], dataset: str, version: int, source_id: int) -> dict[str, Any]:
    """Backward-compatible one-sheet entry point."""
    return export_many_copy(source_path, output_path, snapshot, [{
        "sheet_name": sheet_name, "header_row": header_row, "row_count": row_count,
        "columns": columns, "rows": rows, "original_rows": original_rows,
        "dataset": dataset, "version": version, "source_id": source_id,
    }])
