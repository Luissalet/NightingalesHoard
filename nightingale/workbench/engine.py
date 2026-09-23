"""The DuckDB engine: one persistent read-write connection, table materialisation,
profiling and ingestion (CSV/TSV incl. Spanish formats, XLSX/XLS, Parquet,
JSON/NDJSON, SQLite, folder glob, URL, pasted text).

Ingestion (encoding fallback, day-first date sniffing, Spanish decimal/
thousands detection) is adapted from Laplace's Hoard's `engines/data.py`
(same author, MIT) — comments below mark exactly what was reused. Nightingale
differs in one important way: Laplace's catalogue is read-only (a query
surface over files); this workbench *mutates* data, so every dataset version
is a real materialised `TABLE` (never a lazy view), named `ds_{id}_v{n}`, and
the connection stays read-write for the app's whole lifetime under one lock
instead of switching between a read-only/read-write pair.
"""

from __future__ import annotations

import csv
import io
import json
import re
import sqlite3
import threading
import time
import unicodedata
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, time as dtime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterator, Optional
from uuid import UUID

import duckdb
import openpyxl

from .steps import StepContext, build_step_sql

__all__ = ["Engine", "DataError", "slugify_name", "json_safe"]

DEFAULT_QUERY_TIMEOUT_S = 30
MAX_QUERY_LIMIT = 5000
MAX_CELL_CHARS = 500
MAX_INGEST_PREVIEW_ROWS = 5

_NUMERIC_TYPES = {
    "TINYINT", "SMALLINT", "INTEGER", "BIGINT", "HUGEINT", "UTINYINT",
    "USMALLINT", "UINTEGER", "UBIGINT", "UHUGEINT", "FLOAT", "DOUBLE",
    "REAL", "DECIMAL", "NUMERIC",
}
_TEMPORAL_PREFIXES = ("DATE", "TIMESTAMP", "TIME")

# adapted from Laplace's Hoard: DuckDB's CSV reader accepts these encoding
# names without needing an extra extension; 'latin-1' stands in for
# Windows-1252 (identical for accented Spanish/French/German letters).
_CSV_ENCODINGS = {"utf-8", "utf8", "utf-16", "latin-1", "latin1"}
_CSV_ENCODING_FALLBACKS = ("latin-1",)


class DataError(ValueError):
    pass


def _esc(text: str) -> str:
    return str(text).replace("'", "''")


def q(name: str) -> str:
    return '"' + str(name).replace('"', '""') + '"'


def slugify_name(name: str) -> str:
    """Turn a file stem / sheet / table name into a safe SQL identifier.

    Adapted from Laplace's Hoard's `slugify_name`.
    """
    text = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9_]+", "_", text)
    text = re.sub(r"_{3,}", "__", text).strip("_")
    if not text:
        text = "dataset"
    if text[0].isdigit():
        text = "t_" + text
    return text[:63].lower()


def _base_type(type_name: str) -> str:
    return type_name.upper().split("(")[0].strip()


def is_numeric_type(type_name: str) -> bool:
    return _base_type(type_name) in _NUMERIC_TYPES


def is_temporal_type(type_name: str) -> bool:
    return _base_type(type_name).startswith(_TEMPORAL_PREFIXES)


def _is_nested_type(type_name: str) -> bool:
    t = type_name.upper().rstrip()
    return _base_type(t) in ("STRUCT", "MAP", "UNION") or t.endswith("]")


def json_safe(value: Any, *, cap: Optional[int] = None) -> Any:
    """Make a DuckDB value JSON-safe. Adapted from Laplace's Hoard's `_json_safe`."""
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return str(value)
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date, dtime)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, (bytes, bytearray)):
        return f"<binary, {len(value):,} bytes>"
    if isinstance(value, str):
        if cap is not None and len(value) > cap:
            return value[:cap] + "…"
        return value
    if isinstance(value, dict):
        return {str(k): json_safe(v, cap=cap) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(v, cap=cap) for v in value]
    return str(value)


def _first_lines(text: str, max_chars: int = 500) -> str:
    text = text.strip()
    return text if len(text) <= max_chars else text[:max_chars] + "…"


_ingest_counter = 0


def _temp_table_name() -> str:
    global _ingest_counter
    _ingest_counter += 1
    return f"__ingest_probe_{_ingest_counter}_{int(time.time() * 1000)}"


@dataclass
class IngestResult:
    select_sql: str
    columns: list[dict]
    row_count: int
    extra: dict
    table_name: str = ""


class Engine:
    """One persistent, lock-protected, read-write DuckDB connection."""

    def __init__(self, db_path: Path, cache_dir: Optional[Path] = None):
        self.db_path = Path(db_path)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.cache_dir = Path(cache_dir) if cache_dir else self.db_path.parent / "cache"
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        offline = {"autoinstall_known_extensions": False, "autoload_known_extensions": False}
        self._conn = duckdb.connect(str(self.db_path), config=offline)

    def close(self) -> None:
        with self._lock:
            self._conn.close()

    @contextmanager
    def lock(self) -> Iterator[duckdb.DuckDBPyConnection]:
        with self._lock:
            yield self._conn

    # -- table lifecycle -------------------------------------------------

    def materialize(self, table_name: str, select_sql: str) -> tuple[int, list[dict]]:
        """CREATE TABLE table_name AS select_sql; returns (row_count, columns)."""
        with self._lock:
            for stmt in (f"DROP TABLE IF EXISTS {q(table_name)}", f"DROP VIEW IF EXISTS {q(table_name)}"):
                try:
                    self._conn.execute(stmt)
                except duckdb.Error:
                    pass
            try:
                self._conn.execute(f"CREATE TABLE {q(table_name)} AS {select_sql}")
            except duckdb.Error as exc:
                raise DataError(f"SQL error building this step: {_first_lines(str(exc))}") from exc
            row_count = self._conn.execute(f"SELECT COUNT(*) FROM {q(table_name)}").fetchone()[0]
            columns = [{"name": r[0], "type": r[1]} for r in self._conn.execute(f"DESCRIBE {q(table_name)}").fetchall()]
            return int(row_count), columns

    def materialize_from_dataframe(self, table_name: str, df) -> tuple[int, list[dict]]:
        """Register a pandas DataFrame and materialize it as a table (used to
        write model predictions/labels back as a new dataset version)."""
        with self._lock:
            self._conn.register("__df_in", df)
            try:
                for stmt in (f"DROP TABLE IF EXISTS {q(table_name)}", f"DROP VIEW IF EXISTS {q(table_name)}"):
                    try:
                        self._conn.execute(stmt)
                    except duckdb.Error:
                        pass
                self._conn.execute(f"CREATE TABLE {q(table_name)} AS SELECT * FROM __df_in")
            finally:
                self._conn.unregister("__df_in")
            row_count = self._conn.execute(f"SELECT COUNT(*) FROM {q(table_name)}").fetchone()[0]
            columns = [{"name": r[0], "type": r[1]} for r in self._conn.execute(f"DESCRIBE {q(table_name)}").fetchall()]
            return int(row_count), columns

    def set_view(self, view_name: str, table_name: str) -> None:
        """Point the friendly view name (the dataset's slug) at its current version's table,
        so `data_query`/quality/ask can address a dataset by name instead of an internal id."""
        with self._lock:
            self._conn.execute(f"CREATE OR REPLACE VIEW {q(view_name)} AS SELECT * FROM {q(table_name)}")

    def drop_view(self, view_name: str) -> None:
        with self._lock:
            try:
                self._conn.execute(f"DROP VIEW IF EXISTS {q(view_name)}")
            except duckdb.Error:
                pass

    def to_dataframe(self, table_name: str, limit: Optional[int] = None):
        import pandas as pd  # local import: only model training needs pandas

        sql = f"SELECT * FROM {q(table_name)}" + (f" LIMIT {int(limit)}" if limit else "")
        with self._lock:
            return self._conn.execute(sql).fetch_df()

    def drop_table(self, table_name: str) -> None:
        with self._lock:
            for stmt in (f"DROP TABLE IF EXISTS {q(table_name)}", f"DROP VIEW IF EXISTS {q(table_name)}"):
                try:
                    self._conn.execute(stmt)
                except duckdb.Error:
                    pass

    def rename_table(self, old_name: str, new_name: str) -> None:
        """Move an already-materialized table to its final name (no data is re-scanned).
        Used after ingestion: the probe table that was scanned from the source file/URL/paste
        becomes the dataset's real version-0 table without re-running the source query — that
        query may reference a temp file that ingestion has already deleted by this point."""
        with self._lock:
            if old_name == new_name:
                return
            for stmt in (f"DROP TABLE IF EXISTS {q(new_name)}", f"DROP VIEW IF EXISTS {q(new_name)}"):
                try:
                    self._conn.execute(stmt)
                except duckdb.Error:
                    pass
            self._conn.execute(f"ALTER TABLE {q(old_name)} RENAME TO {q(new_name)}")

    def table_columns(self, table_name: str) -> list[str]:
        with self._lock:
            return [r[0] for r in self._conn.execute(f"DESCRIBE {q(table_name)}").fetchall()]

    def row_count(self, table_name: str) -> int:
        with self._lock:
            return int(self._conn.execute(f"SELECT COUNT(*) FROM {q(table_name)}").fetchone()[0])

    # -- profiling ---------------------------------------------------------

    def profile_table(self, table_name: str, columns: list[dict], total: Optional[int] = None) -> dict:
        """Column profiles: type, nulls %, distinct, min/max/mean/sd, histogram,
        top values, outliers by IQR. Adapted from Laplace's Hoard's `_profile`,
        extended with an IQR outlier count for numeric columns."""
        with self._lock:
            conn = self._conn
            if total is None:
                total = int(conn.execute(f"SELECT COUNT(*) FROM {q(table_name)}").fetchone()[0])
            profile: dict[str, Any] = {}
            for col in columns:
                cname = col["name"]
                c = q(cname)
                nulls, distinct = conn.execute(
                    f"SELECT COUNT(*) - COUNT({c}), APPROX_COUNT_DISTINCT({c}) FROM {q(table_name)}"
                ).fetchone()
                entry: dict[str, Any] = {
                    "type": col["type"],
                    "nulls": int(nulls or 0),
                    "nulls_pct": round(100.0 * nulls / total, 2) if total else 0.0,
                    "distinct_approx": int(distinct) if distinct is not None else 0,
                }
                if is_numeric_type(col["type"]):
                    row = conn.execute(
                        f"SELECT MIN({c}), MAX({c}), AVG({c}), STDDEV_SAMP({c}), "
                        f"QUANTILE_CONT({c}, 0.25), QUANTILE_CONT({c}, 0.5), QUANTILE_CONT({c}, 0.75) "
                        f"FROM {q(table_name)}"
                    ).fetchone()
                    lo, hi, mean, sd, p25, median, p75 = row
                    entry.update({"min": json_safe(lo), "max": json_safe(hi), "mean": json_safe(mean),
                                  "sd": json_safe(sd), "median": json_safe(median)})
                    if p25 is not None and p75 is not None:
                        # DECIMAL columns come back as decimal.Decimal, which can't mix with a
                        # plain float in arithmetic — normalize to float before the fence math.
                        p25_f, p75_f = float(p25), float(p75)
                        iqr = p75_f - p25_f
                        fence_lo, fence_hi = p25_f - 1.5 * iqr, p75_f + 1.5 * iqr
                        outliers = conn.execute(
                            f"SELECT COUNT(*) FROM {q(table_name)} WHERE {c} < {fence_lo} OR {c} > {fence_hi}"
                        ).fetchone()[0]
                        entry["outliers_iqr"] = int(outliers)
                    try:
                        hist = conn.execute(
                            f"SELECT bin, COUNT(*) FROM (SELECT LEAST(9, FLOOR(({c} - m.lo) / "
                            f"NULLIF(m.hi - m.lo, 0) * 10))::INT AS bin FROM {q(table_name)}, "
                            f"(SELECT MIN({c}) AS lo, MAX({c}) AS hi FROM {q(table_name)}) m "
                            f"WHERE {c} IS NOT NULL) GROUP BY bin ORDER BY bin"
                        ).fetchall()
                        counts = [0] * 10
                        for b, n in hist:
                            counts[int(b) if b is not None else 0] += int(n)
                        entry["histogram"] = counts
                    except duckdb.Error:
                        pass
                elif is_temporal_type(col["type"]):
                    row = conn.execute(f"SELECT MIN({c}), MAX({c}) FROM {q(table_name)}").fetchone()
                    entry.update({"min": json_safe(row[0]), "max": json_safe(row[1])})
                elif _is_nested_type(col["type"]):
                    entry["note"] = "nested value"
                else:
                    try:
                        top = conn.execute(
                            f"SELECT {c}, COUNT(*) AS n FROM {q(table_name)} WHERE {c} IS NOT NULL "
                            f"GROUP BY {c} ORDER BY n DESC, 1 LIMIT 5"
                        ).fetchall()
                        entry["top_values"] = [{"value": json_safe(v, cap=80), "count": int(n)} for v, n in top]
                    except duckdb.Error:
                        entry["top_values"] = []
                profile[cname] = entry
            return profile

    def correlation_matrix(self, table_name: str, columns: list[str]) -> dict:
        with self._lock:
            conn = self._conn
            n = len(columns)
            matrix = [[None] * n for _ in range(n)]
            for i, a in enumerate(columns):
                for j, b in enumerate(columns):
                    if j < i:
                        matrix[i][j] = matrix[j][i]
                        continue
                    if a == b:
                        matrix[i][j] = 1.0
                        continue
                    try:
                        val = conn.execute(f"SELECT CORR({q(a)}, {q(b)}) FROM {q(table_name)}").fetchone()[0]
                    except duckdb.Error:
                        val = None
                    matrix[i][j] = round(val, 4) if val is not None else None
            return {"columns": columns, "matrix": matrix}

    def sample_rows(self, table_name: str, columns: list[str], limit: int = 5) -> list[dict]:
        with self._lock:
            rows = self._conn.execute(f"SELECT * FROM {q(table_name)} LIMIT {int(limit)}").fetchall()
        return [{columns[i]: json_safe(v, cap=120) for i, v in enumerate(row)} for row in rows]

    # -- steps -------------------------------------------------------------

    def apply_step(self, prev_table: str, new_table: str, op: str, params: dict,
                    resolve_dataset_table) -> tuple[str, int, list[dict]]:
        columns = self.table_columns(prev_table)
        ctx = StepContext(prev_table=prev_table, columns=columns, resolve_dataset_table=resolve_dataset_table)
        select_sql = build_step_sql(op, ctx, params)
        row_count, out_columns = self.materialize(new_table, select_sql)
        return select_sql, row_count, out_columns

    def preview_step(self, prev_table: str, op: str, params: dict, resolve_dataset_table,
                      n: int = 20) -> dict:
        columns = self.table_columns(prev_table)
        prev_count = self.row_count(prev_table)
        ctx = StepContext(prev_table=prev_table, columns=columns, resolve_dataset_table=resolve_dataset_table)
        select_sql = build_step_sql(op, ctx, params)
        with self._lock:
            try:
                described = [{"name": r[0], "type": r[1]} for r in self._conn.execute(f"DESCRIBE ({select_sql})").fetchall()]
                rows = self._conn.execute(f"SELECT * FROM ({select_sql}) LIMIT {int(n)}").fetchall()
                new_count = self._conn.execute(f"SELECT COUNT(*) FROM ({select_sql})").fetchone()[0]
            except duckdb.Error as exc:
                raise DataError(f"SQL error building this step: {_first_lines(str(exc))}") from exc
        out_cols = [c["name"] for c in described]
        preview_rows = [{out_cols[i]: json_safe(v, cap=200) for i, v in enumerate(row)} for row in rows]
        before_names, after_names = set(columns), {c["name"] for c in described}
        return {
            "select_sql": select_sql,
            "columns_before": columns,
            "columns_after": described,
            "columns_added": sorted(after_names - before_names),
            "columns_removed": sorted(before_names - after_names),
            "row_count_before": prev_count,
            "row_count_after": int(new_count),
            "row_count_delta": int(new_count) - prev_count,
            "preview_rows": preview_rows,
        }

    def join_match_rate(self, left_table: str, right_table: str, on: list[dict]) -> dict:
        conds = " AND ".join(f"l.{q(p['left'])} = r.{q(p['right'])}" for p in on)
        with self._lock:
            left_total = self._conn.execute(f"SELECT COUNT(*) FROM {q(left_table)}").fetchone()[0]
            matched = self._conn.execute(
                f"SELECT COUNT(*) FROM {q(left_table)} l WHERE EXISTS (SELECT 1 FROM {q(right_table)} r WHERE {conds})"
            ).fetchone()[0]
        return {"left_rows": int(left_total), "matched": int(matched),
                "unmatched": int(left_total) - int(matched),
                "match_rate": round(matched / left_total, 4) if left_total else 0.0}

    # -- query ---------------------------------------------------------------

    def query(self, sql: str, limit: int = 200, timeout_s: int = DEFAULT_QUERY_TIMEOUT_S,
              cell_cap: Optional[int] = MAX_CELL_CHARS) -> dict:
        limit = max(1, min(int(limit), MAX_QUERY_LIMIT))
        with self._lock:
            conn = self._conn
            timer = threading.Timer(timeout_s, conn.interrupt)
            timer.daemon = True
            start = time.monotonic()
            timer.start()
            try:
                rel = conn.execute(sql)
                if rel.description is None:
                    raise DataError("the statement returned no result set")
                columns = [{"name": d[0], "type": str(d[1])} for d in rel.description]
                rows = rel.fetchmany(limit + 1)
                del rel
            except duckdb.InterruptException as exc:
                raise DataError(f"query exceeded {timeout_s}s and was stopped") from exc
            except duckdb.Error as exc:
                raise DataError(f"SQL error: {_first_lines(str(exc))}") from exc
            finally:
                timer.cancel()
            elapsed_ms = (time.monotonic() - start) * 1000
        truncated = len(rows) > limit
        rows = rows[:limit]
        col_names = [c["name"] for c in columns]
        json_rows = [{col_names[i]: json_safe(v, cap=cell_cap) for i, v in enumerate(row)} for row in rows]
        return {"columns": columns, "rows": json_rows, "row_count": len(json_rows),
                "truncated": truncated, "elapsed_ms": round(elapsed_ms, 2)}

    # -- ingestion -----------------------------------------------------------
    # CSV/date/number handling below is adapted from Laplace's Hoard's
    # `engines/data.py` (`_register_csv`, `_sniff_dayfirst_format`,
    # `_apply_locale_numbers`, `_resolve_excel_skip_rows`) — same author, MIT.

    def build_csv_select(self, path: Path, options: dict) -> tuple[str, dict]:
        delim = options.get("delimiter")
        header = options.get("header", True)
        encoding = options.get("encoding")
        date_format = options.get("date_format")
        if encoding and str(encoding).lower() not in _CSV_ENCODINGS:
            raise DataError(f"unsupported encoding: {encoding!r}; use utf-8, utf-16 or latin-1")
        date_fmt = date_format or _sniff_dayfirst_format(path, delim, bool(header), str(encoding or "utf-8"))

        def build(enc: Optional[str]) -> str:
            args = [f"'{_esc(str(path))}'", f"header={str(bool(header)).lower()}"]
            if delim:
                args.append(f"delim='{_esc(str(delim))}'")
            elif path.suffix.lower() == ".tsv":
                args.append("delim='\\t'")
            if enc:
                args.append(f"encoding='{_esc(str(enc))}'")
            if date_fmt:
                args.append(f"dateformat='{_esc(date_fmt)}'")
            return f"SELECT * FROM read_csv_auto({', '.join(args)})"

        candidates = [encoding] if encoding else [None, *_CSV_ENCODING_FALLBACKS]
        last_exc = None
        for enc in candidates:
            select_sql = build(enc)
            try:
                with self._lock:
                    self._conn.execute(f"SELECT COUNT(*) FROM ({select_sql}) LIMIT 0")
                return select_sql, {"encoding_detected": enc} if (enc and not encoding) else {}
            except duckdb.Error as exc:
                if encoding or "not utf-8 encoded" not in str(exc).lower() and "invalid unicode" not in str(exc).lower():
                    raise DataError(f"could not read {path.name}: {_first_lines(str(exc))}") from exc
                last_exc = exc
                continue
        raise DataError(
            f"could not read {path.name}: not valid UTF-8 and latin-1 did not work either; "
            f"pass options.encoding explicitly ({_first_lines(str(last_exc))})"
        ) from last_exc

    def ingest_delimited(self, path: Path, options: dict, source_col: Optional[str] = None) -> IngestResult:
        select_sql, extra = self.build_csv_select(path, options)
        decimal_sep, thousands_sep = options.get("decimal_separator"), options.get("thousands_separator")
        with self._lock:
            select_sql, converted = _apply_locale_numbers(
                self._conn, select_sql, decimal_sep, thousands_sep, auto=not (decimal_sep or thousands_sep)
            )
        if source_col:
            select_sql = f"SELECT *, '{_esc(str(path))}' AS {q(source_col)} FROM ({select_sql}) __t"
        tmp = _temp_table_name()
        row_count, columns = self.materialize(tmp, select_sql)
        if converted:
            extra["numbers_converted"] = converted
        return IngestResult(select_sql=select_sql, columns=columns, row_count=row_count, extra=extra, table_name=tmp)

    def ingest_parquet(self, path: Path) -> IngestResult:
        select_sql = f"SELECT * FROM read_parquet('{_esc(str(path))}')"
        tmp = _temp_table_name()
        row_count, columns = self.materialize(tmp, select_sql)
        return IngestResult(select_sql=select_sql, columns=columns, row_count=row_count, extra={}, table_name=tmp)

    def ingest_folder(self, folder: Path, glob: str, source_col: Optional[str] = "_source_file") -> IngestResult:
        if "/" in glob or "\\" in glob or ".." in glob:
            raise DataError("glob must be a file pattern inside the folder, e.g. '*.csv'")
        files = [f for f in folder.glob(glob) if f.is_file()]
        if not files:
            raise DataError(f"no files match {glob!r} in {folder}")
        glob_path = str(folder / glob)
        fn = "read_parquet" if glob.endswith(".parquet") else "read_csv_auto"
        filename_arg = ", filename=true" if fn == "read_csv_auto" else ""
        select_sql = f"SELECT * FROM {fn}('{_esc(glob_path)}'{filename_arg})"
        if source_col and fn == "read_csv_auto":
            select_sql = f"SELECT * EXCLUDE (filename), filename AS {q(source_col)} FROM ({select_sql}) __t"
        tmp = _temp_table_name()
        row_count, columns = self.materialize(tmp, select_sql)
        return IngestResult(select_sql=select_sql, columns=columns, row_count=row_count,
                             extra={"files": len(files)}, table_name=tmp)

    def ingest_json(self, path: Path, options: dict) -> IngestResult:
        flatten = bool(options.get("flatten"))
        select_sql = f"SELECT * FROM read_json_auto('{_esc(str(path))}')"
        probe = _temp_table_name()
        row_count, columns = self.materialize(probe, select_sql)
        if flatten:
            nested = [c["name"] for c in columns if _is_nested_type(c["type"])]
            if nested:
                kept = [c["name"] for c in columns if c["name"] not in nested]
                kept_sql = ", ".join(q(c) for c in kept)
                unnest_parts = ", ".join(f"UNNEST({q(c)}, max_depth := 2)" for c in nested)
                candidate_sql = f"SELECT {kept_sql}{', ' if kept else ''}{unnest_parts} FROM {q(probe)}"
                with self._lock:
                    try:
                        self._conn.execute(f"DESCRIBE ({candidate_sql})")
                        select_sql = candidate_sql
                    except duckdb.Error:
                        pass  # flattening not possible for this shape: keep the nested column as-is
        tmp = probe
        if select_sql != f"SELECT * FROM read_json_auto('{_esc(str(path))}')":
            tmp = _temp_table_name()
            row_count, columns = self.materialize(tmp, select_sql)
            self.drop_table(probe)
        return IngestResult(select_sql=select_sql, columns=columns, row_count=row_count,
                             extra={"flattened": flatten}, table_name=tmp)

    def ingest_pasted_text(self, text: str, fmt: str, options: dict) -> IngestResult:
        suffix = ".json" if fmt == "json" else ".csv"
        tmp_path = self.cache_dir / f"__paste_{int(time.time() * 1000)}{suffix}"
        tmp_path.write_text(text, encoding="utf-8")
        try:
            if fmt == "json":
                return self.ingest_json(tmp_path, options)
            return self.ingest_delimited(tmp_path, options)
        finally:
            tmp_path.unlink(missing_ok=True)  # the data is already materialized in DuckDB by now

    def ingest_excel(self, path: Path, options: dict) -> list[tuple[str, IngestResult]]:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            wanted = options.get("sheets") or ([options["sheet"]] if options.get("sheet") else None)
            sheets = wanted or wb.sheetnames
            missing = [s for s in sheets if s not in wb.sheetnames]
            if missing:
                raise DataError(f"sheet(s) not found: {missing}; workbook has {wb.sheetnames}")
            explicit_skip = options.get("skip_rows")
            out = []
            for sheet_name in sheets:
                ws = wb[sheet_name]
                skip_rows = _resolve_excel_skip_rows(ws, explicit_skip)
                rows_iter = ws.iter_rows(min_row=skip_rows + 1, values_only=True)
                try:
                    header = next(rows_iter)
                except StopIteration:
                    continue
                header = [str(h) if h is not None else f"col{i}" for i, h in enumerate(header)]
                csv_path = self.cache_dir / f"__xlsx_{slugify_name(sheet_name)}_{int(time.time()*1000)}.csv"
                with csv_path.open("w", newline="", encoding="utf-8") as fh:
                    writer = csv.writer(fh)
                    writer.writerow(header)
                    for row in rows_iter:
                        writer.writerow(["" if v is None else v for v in row])
                select_sql = f"SELECT * FROM read_csv_auto('{_esc(str(csv_path))}', header=true)"
                with self._lock:
                    select_sql, converted = _apply_locale_numbers(self._conn, select_sql, None, None, auto=True)
                tmp = _temp_table_name()
                row_count, columns = self.materialize(tmp, select_sql)
                csv_path.unlink(missing_ok=True)
                extra = {"sheet": sheet_name}
                if explicit_skip is None and skip_rows:
                    extra["skip_rows_detected"] = skip_rows
                if converted:
                    extra["numbers_converted"] = converted
                out.append((sheet_name, IngestResult(select_sql=select_sql, columns=columns, row_count=row_count,
                                                       extra=extra, table_name=tmp)))
            if not out:
                raise DataError("workbook has no readable sheets")
            return out
        finally:
            wb.close()

    def ingest_sqlite(self, path: Path, options: dict) -> list[tuple[str, IngestResult]]:
        src = sqlite3.connect(f"{path.as_uri()}?mode=ro", uri=True)
        try:
            all_tables = [r[0] for r in src.execute(
                "SELECT name FROM sqlite_master WHERE type IN ('table','view') AND name NOT LIKE 'sqlite_%'"
            ).fetchall()]
            tables = options.get("tables") or all_tables
            missing = [t for t in tables if t not in all_tables]
            if missing:
                raise DataError(f"table(s) not found: {missing}; database has {all_tables}")
            out = []
            for table in tables:
                quoted = '"' + table.replace('"', '""') + '"'
                cur = src.execute(f"SELECT * FROM {quoted}")
                cols = [d[0] for d in cur.description]
                csv_path = self.cache_dir / f"__sqlite_{slugify_name(table)}_{int(time.time()*1000)}.csv"
                with csv_path.open("w", newline="", encoding="utf-8") as fh:
                    writer = csv.writer(fh)
                    writer.writerow(cols)
                    for row in cur:
                        writer.writerow(list(row))
                select_sql = f"SELECT * FROM read_csv_auto('{_esc(str(csv_path))}', header=true)"
                tmp = _temp_table_name()
                row_count, columns = self.materialize(tmp, select_sql)
                csv_path.unlink(missing_ok=True)
                out.append((table, IngestResult(select_sql=select_sql, columns=columns, row_count=row_count,
                                                  extra={"table": table}, table_name=tmp)))
            if not out:
                raise DataError("sqlite file has no tables")
            return out
        finally:
            src.close()


# ---------------------------------------------------------------------------
# CSV locale handling — adapted from Laplace's Hoard's `engines/data.py`.
# ---------------------------------------------------------------------------

_DATE_TOKEN_RE = re.compile(r"^(\d{1,2})([/-])(\d{1,2})\2(\d{2}|\d{4})$")


def _sniff_dayfirst_format(p: Path, delimiter: Optional[str], header: bool, encoding: str = "utf-8") -> Optional[str]:
    sep_guess = delimiter or ("\t" if p.suffix.lower() == ".tsv" else ",")
    dayfirst: Optional[bool] = None
    sep = "/"
    year_len4 = True
    try:
        with p.open("r", encoding=encoding, errors="replace", newline="") as fh:
            reader = csv.reader(fh, delimiter=sep_guess)
            for i, row in enumerate(reader):
                if i == 0 and header:
                    continue
                if i > 500:
                    break
                for cell in row:
                    m = _DATE_TOKEN_RE.match((cell or "").strip())
                    if not m:
                        continue
                    a_s, s, b_s, y = m.groups()
                    a, b = int(a_s), int(b_s)
                    if a > 31 or b > 31 or a == 0 or b == 0:
                        continue
                    if a > 12 and b <= 12:
                        if dayfirst is False:
                            return None
                        dayfirst, sep, year_len4 = True, s, len(y) == 4
                    elif b > 12 and a <= 12:
                        if dayfirst is True:
                            return None
                        dayfirst, sep, year_len4 = False, s, len(y) == 4
    except OSError:
        return None
    if dayfirst:
        return f"%d{sep}%m{sep}" + ("%Y" if year_len4 else "%y")
    return None


_MAX_TITLE_ROWS = 10


def _resolve_excel_skip_rows(ws, explicit: Any) -> int:
    if explicit is not None:
        try:
            return max(0, int(explicit))
        except (TypeError, ValueError):
            raise DataError(f"skip_rows must be a whole number, got {explicit!r}")
    try:
        peek = list(ws.iter_rows(min_row=1, max_row=_MAX_TITLE_ROWS + 1, values_only=True))
    except Exception:  # noqa: BLE001
        return 0
    nonempty = lambda row: sum(1 for v in row if v not in (None, "") and str(v).strip())  # noqa: E731
    if not peek or nonempty(peek[0]) != 1:
        return 0
    for i, row in enumerate(peek[1:], start=1):
        if nonempty(row) >= 2:
            return i
    return 0


_COMMA_NUMBER_RE = r"[+-]?(\d{1,3}(\.\d{3})+|\d+)(,\d+)?"
_COMMA_DECIMAL_EVIDENCE_RE = r"[+-]?(\d+,(\d{1,2}|\d{4,})|\d{1,3}(\.\d{3})+,\d+)"
_MAX_EXACT_SCALE = 6


def _apply_locale_numbers(conn, base_sql: str, decimal_sep: Optional[str], thousands_sep: Optional[str],
                           *, auto: bool = False) -> tuple[str, list[str]]:
    try:
        described = conn.execute(f"DESCRIBE ({base_sql})").fetchall()
    except duckdb.Error:
        return base_sql, []
    if auto:
        decimal_sep, thousands_sep = ",", "."
    decimal_sep = decimal_sep or "."
    parts: list[str] = []
    converted: list[str] = []
    for row in described:
        cname, ctype = row[0], row[1]
        col = q(cname)
        if _base_type(ctype) != "VARCHAR" or (decimal_sep == "." and not thousands_sep):
            parts.append(col)
            continue
        norm = f"TRIM({col})"
        if thousands_sep:
            norm = f"REPLACE({norm}, '{_esc(thousands_sep)}', '')"
        if decimal_sep != ".":
            norm = f"REPLACE({norm}, '{_esc(decimal_sep)}', '.')"
        blank = f"({col} IS NULL OR TRIM({col}) = '')"
        try:
            if auto:
                non_blank, matching, evidence, scale = conn.execute(
                    f"SELECT COUNT(*) FILTER (WHERE NOT {blank}), "
                    f"COUNT(*) FILTER (WHERE NOT {blank} AND regexp_full_match(TRIM({col}), '{_COMMA_NUMBER_RE}')), "
                    f"COUNT(*) FILTER (WHERE NOT {blank} AND regexp_full_match(TRIM({col}), '{_COMMA_DECIMAL_EVIDENCE_RE}')), "
                    f"MAX(CASE WHEN strpos(TRIM({col}), ',') > 0 THEN length(split_part(TRIM({col}), ',', 2)) ELSE 0 END) "
                    f"FROM ({base_sql})"
                ).fetchone()
                ok = bool(non_blank) and matching == non_blank and evidence > 0
                target = f"DECIMAL(18, {int(scale or 0)})" if (scale or 0) <= _MAX_EXACT_SCALE else "DOUBLE"
            else:
                probe = f"TRY_CAST(NULLIF({norm}, '') AS DOUBLE)"
                non_blank, castable = conn.execute(
                    f"SELECT COUNT(*) FILTER (WHERE NOT {blank}), "
                    f"COUNT(*) FILTER (WHERE NOT {blank} AND {probe} IS NOT NULL) FROM ({base_sql})"
                ).fetchone()
                ok = bool(non_blank) and castable / non_blank >= 0.9
                target = "DOUBLE"
        except duckdb.Error:
            parts.append(col)
            continue
        if ok:
            parts.append(f"TRY_CAST(NULLIF({norm}, '') AS {target}) AS {col}")
            converted.append(cname)
        else:
            parts.append(col)
    if not converted:
        return base_sql, []
    return f"SELECT {', '.join(parts)} FROM ({base_sql}) t", converted
