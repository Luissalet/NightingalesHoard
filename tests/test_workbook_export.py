"""End-to-end tests for conservative, workbook-preserving XLSX exports."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import re
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

from nightingale.agent_tools import tool_catalog
from nightingale.workbook_export import _cell_xml, _excel_value, _patch_worksheet_cells, inventory


def _rich_workbook(path: Path, *, target_formula: bool = False, target_validation: bool = False,
                   date1904: bool = False, include_inventory: bool = False) -> Path:
    wb = openpyxl.Workbook()
    if date1904:
        wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
    data = wb.active
    data.title = "Ventas"
    data.append(["fecha", "importe", "categoria", "activo", "literal", "ref_text", "fecha_text"])
    data.append([openpyxl.utils.datetime.from_ISO8601("2026-01-03T00:00:00"), 12.5, "café", True, "=texto", "00123", "2026-01-03"])
    data.append([openpyxl.utils.datetime.from_ISO8601("2026-01-04T00:00:00"), 24.75, "té", False, "=literal", "00124", "2026-01-04"])
    data.append([openpyxl.utils.datetime.from_ISO8601("2026-01-05T00:00:00"), 8.0, "niño", True, "=sin fórmula", "00125", "2026-01-05"])
    for row in range(2, 5):
        data.cell(row, 5).data_type = "s"
    data["A2"].number_format = "dd/mm/yyyy"
    data["A3"].number_format = "dd/mm/yyyy"
    data["A4"].number_format = "dd/mm/yyyy"
    for cell in data[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="315A72")
    for cell in data["B"][1:]:
        cell.number_format = '#,##0.00 [$€-es-ES]'
    table = Table(displayName="SalesTable", ref="A1:G4")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    data.add_table(table)
    data.auto_filter.ref = "A1:G4"
    if target_formula:
        data["B2"] = "=6+6.5"
    if target_validation:
        from openpyxl.worksheet.datavalidation import DataValidation
        validation = DataValidation(type="list", formula1='"café,té,niño"', allow_blank=False)
        data.add_data_validation(validation)
        validation.add("C2:C4")
        data["C2"].comment = openpyxl.comments.Comment("Keep this note while editing.", "Nightingale QA")
        data["C2"].hyperlink = "https://example.invalid/category"

    if include_inventory:
        inventory_sheet = wb.create_sheet("Inventario")
        inventory_sheet.append(["sku", "stock", "categoria"])
        inventory_sheet.append(["A-01", 4, "inicial"])
        inventory_sheet.append(["B-02", 9, "inicial"])
        inventory_sheet.append(["C-03", 2, "reserva"])
        inventory_table = Table(displayName="InventoryTable", ref="A1:C4")
        inventory_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        inventory_sheet.add_table(inventory_table)
    summary = wb.create_sheet("Resumen")
    summary["A1"] = "Total de ventas"
    summary["B1"] = "=SUM(Ventas!B2:B4)"
    summary["B1"].number_format = '#,##0.00 [$€-es-ES]'
    summary["A2"] = "Revisar fuente"
    summary["A2"].hyperlink = "https://example.invalid/report"
    summary["A2"].style = "Hyperlink"
    summary["A3"] = "Comentario conservado"
    summary["A3"].comment = openpyxl.comments.Comment("Nota de prueba: revisión manual.", "Nightingale QA")

    notes = wb.create_sheet("Notas")
    notes.append(["Estado", "Valor"])
    notes.append(["Aprobado", 2])
    notes.append(["Pendiente", 4])
    notes["A2"].comment = openpyxl.comments.Comment("Comentario no afectado.", "Nightingale QA")
    notes["A3"].hyperlink = "https://example.invalid/estado"
    notes["A3"].style = "Hyperlink"
    notes["B2"].fill = openpyxl.styles.PatternFill("solid", fgColor="FCE4D6")
    notes["A2"].data_type = "s"
    from openpyxl.worksheet.datavalidation import DataValidation
    validation = DataValidation(type="list", formula1='"Aprobado,Pendiente"', allow_blank=False)
    notes.add_data_validation(validation)
    validation.add("A2:A3")
    chart = BarChart()
    chart.title = "Conteos de prueba"
    chart.add_data(Reference(notes, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(notes, min_col=1, min_row=2, max_row=3))
    notes.add_chart(chart, "D2")
    wb.save(path)
    # Add a real Office ignorable namespace attribute to worksheet rows and a
    # stale formula cache. The exporter must preserve these untouched XML bytes.
    patched = path.with_suffix(".fixture-tmp.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(patched, "w") as destination:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = payload.decode("utf-8")
                xml = xml.replace('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
                                  '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                                  'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
                                  'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac" mc:Ignorable="x14ac"', 1)
                xml = re.sub(r'<row r="2"', '<row r="2" x14ac:dyDescent="0.25"', xml, count=1)
                payload = xml.encode("utf-8")
            elif info.filename == "xl/worksheets/sheet2.xml":
                payload = payload.replace(b"<v />", b"<v>45.25</v>", 1)
            destination.writestr(info, payload)
    patched.replace(path)
    return path


def _prefix_workbook_xml(path: Path, *, omit_calc_pr: bool) -> None:
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    temp = path.with_suffix(".prefixed.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w") as destination:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename in {"xl/workbook.xml", "xl/worksheets/sheet1.xml"}:
                xml = payload.decode("utf-8")
                xml = xml.replace(f'xmlns="{main_ns}"', f'xmlns:x="{main_ns}"', 1)
                xml = re.sub(r"<(/?)([A-Za-z_][\w.-]*)(?=[\s>])", r"<\1x:\2", xml)
                if info.filename == "xl/workbook.xml" and omit_calc_pr:
                    xml = re.sub(r"<x:calcPr\b[^>]*/>", "", xml)
                payload = xml.encode("utf-8")
            destination.writestr(info, payload)
    temp.replace(path)


def _ingest_target(services, path):
    result = services.ingest_file(str(path), "libro", {"sheets": ["Ventas"]})
    return result["name"]


def _cell_fragment(xml: bytes, ref: str) -> bytes:
    pattern = re.compile(rb"<c\b(?=[^>]*\br=['\"]" + ref.encode() + rb"['\"])[^>]*(?:/>|>.*?</c>)", re.DOTALL)
    match = pattern.search(xml)
    assert match is not None, ref
    return match.group(0)


def test_preserving_export_changes_only_target_sheet_and_keeps_package_features(services, tmp_path):
    source = _rich_workbook(tmp_path / "ventas.xlsx")
    original_bytes = source.read_bytes()
    original_sha = hashlib.sha256(original_bytes).hexdigest()
    source_inventory = inventory(source)
    assert source_inventory["sheets"][0]["tables"] == [{"name": "SalesTable", "ref": "A1:G4", "part": "xl/tables/table1.xml"}]
    assert len(source_inventory["sheets"][1]["comments"]) == 1
    assert source_inventory["sheets"][1]["comments"][0]["refs"] == ["A3"]
    dataset = _ingest_target(services, source)
    assert services.preview(dataset, limit=4)["rows"][0]["categoria"] == "café"
    services.transform_apply(dataset, "derive", {"name": "importe", "expr": "importe + 10"})
    services.transform_apply(dataset, "replace", {"column": "categoria", "pattern": "café", "replacement": "té"})
    changed = services.export(dataset, "xlsx", mode="preserve_workbook", version=2)
    output = Path(changed["path"])
    assert output.exists() and Path(changed["receipt_path"]).exists()
    receipt = json.loads(Path(changed["receipt_path"]).read_text(encoding="utf-8"))
    assert receipt["analysis_log_id"] == changed["log_id"]
    assert receipt["source_sha256"] == original_sha
    assert receipt["dataset_version"] == 2 and receipt["sheet"] == "Ventas"
    assert receipt["data_range"] == "A2:G4"
    assert set(receipt["changed_package_parts"]) == {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
    assert source.read_bytes() == original_bytes and hashlib.sha256(source.read_bytes()).hexdigest() == original_sha

    after = openpyxl.load_workbook(output, data_only=False)
    assert after.sheetnames == ["Ventas", "Resumen", "Notas"]
    assert [after["Ventas"].cell(row, 3).value for row in range(2, 5)] == ["té", "té", "niño"]
    assert [after["Ventas"].cell(row, 2).value for row in range(2, 5)] == [22.5, 34.75, 18]
    assert [after["Ventas"].cell(row, 4).value for row in range(2, 5)] == [True, False, True]
    assert [after["Ventas"].cell(row, 5).value for row in range(2, 5)] == ["=texto", "=literal", "=sin fórmula"]
    assert [after["Ventas"].cell(row, 5).data_type for row in range(2, 5)] == ["s", "s", "s"]
    assert [after["Ventas"].cell(row, 6).value for row in range(2, 5)] == ["00123", "00124", "00125"]
    assert [after["Ventas"].cell(row, 7).value for row in range(2, 5)] == ["2026-01-03", "2026-01-04", "2026-01-05"]
    assert [after["Ventas"].cell(row, 6).data_type for row in range(2, 5)] == ["s", "s", "s"]
    assert [after["Ventas"].cell(row, 7).data_type for row in range(2, 5)] == ["s", "s", "s"]
    assert after["Ventas"]["A2"].number_format == "dd/mm/yyyy"
    assert after["Ventas"]["B2"].number_format == '#,##0.00 [$€-es-ES]'
    assert after["Ventas"].tables["SalesTable"].ref == "A1:G4"
    assert after["Ventas"].auto_filter.ref == "A1:G4"
    assert after["Resumen"]["B1"].value == "=SUM(Ventas!B2:B4)"
    assert after["Resumen"]["A2"].hyperlink.target == "https://example.invalid/report"
    assert after["Resumen"]["A3"].comment.text == "Nota de prueba: revisión manual."
    assert len(after["Notas"]._charts) == 1
    assert str(after["Notas"].data_validations.dataValidation[0].sqref) == "A2:A3"
    assert after["Notas"]["A2"].comment.text == "Comentario no afectado."
    assert after["Notas"]["A3"].hyperlink.target == "https://example.invalid/estado"
    with zipfile.ZipFile(output) as after_zip:
        patched_xml = after_zip.read("xl/worksheets/sheet1.xml")
        assert b'mc:Ignorable="x14ac"' in patched_xml
        assert b'x14ac:dyDescent="0.25"' in patched_xml
        assert patched_xml.count(b'x14ac:dyDescent="0.25"') == 1
        assert after_zip.read("xl/worksheets/sheet2.xml") == zipfile.ZipFile(source).read("xl/worksheets/sheet2.xml")
        with zipfile.ZipFile(source) as before_zip:
            original_xml = before_zip.read("xl/worksheets/sheet1.xml")
        for ref in ("F2", "G2"):
            assert _cell_fragment(patched_xml, ref) == _cell_fragment(original_xml, ref)
    after.close()
    formulas = openpyxl.load_workbook(output, data_only=True)
    assert formulas["Resumen"]["B1"].value == 45.25  # stale cache remains until a spreadsheet engine recalculates
    formulas.close()
    with zipfile.ZipFile(source) as before_zip, zipfile.ZipFile(output) as after_zip:
        for part, digest in source_inventory["package_parts"].items():
            if part not in receipt["changed_package_parts"]:
                assert hashlib.sha256(after_zip.read(part)).hexdigest() == digest, part
        assert b'fullCalcOnLoad="1"' in after_zip.read("xl/workbook.xml")

    # A later request is deterministic in workbook contents even though the log id changes.
    repeated = services.export(dataset, "xlsx", mode="preserve_workbook", version=2, path="ventas-repeat.xlsx")
    assert hashlib.sha256(Path(repeated["path"]).read_bytes()).hexdigest() == changed["output_sha256"]
    assert repeated["source_sha256"] == changed["source_sha256"]
    evidence = os.environ.get("NIGHTINGALE_WORKBOOK_EXPORT_EVIDENCE")
    if evidence:
        evidence_dir = Path(evidence)
        evidence_dir.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, evidence_dir / "source.xlsx")
        shutil.copyfile(output, evidence_dir / "preserved-output.xlsx")
        shutil.copyfile(changed["receipt_path"], evidence_dir / "preserved-output.workbook-export.json")


def test_api_mcp_version_selection_and_undo_are_logged(client, tmp_path):
    source = _rich_workbook(tmp_path / "api-ventas.xlsx")
    dataset = _ingest_target(client.services, source)
    client.services.transform_apply(dataset, "replace", {"column": "categoria", "pattern": "café", "replacement": "té"})
    schema = next(item for item in tool_catalog() if item["name"] == "data_export")["inputSchema"]
    assert schema["properties"]["mode"]["enum"] == ["flat", "preserve_workbook", "inventory"]
    inventory_response = client.get(f"/api/workbooks/{dataset}")
    assert inventory_response.status_code == 200
    inventory_body = inventory_response.json()
    assert inventory_body["sheets"][0]["name"] == "Ventas"
    assert inventory_body["selected_sheet"] == "Ventas" and inventory_body["source_matches_snapshot"]
    inventory_agent = client.post("/api/agent/call", headers={"Authorization": f"Bearer {client.services.token}"},
                                  json={"name": "data_export", "arguments": {"dataset": dataset, "mode": "inventory"}})
    assert inventory_agent.status_code == 200 and inventory_agent.json()["source_matches_snapshot"]
    invalid_inventory_batch = {"dataset": dataset, "format": "xlsx", "mode": "inventory",
                               "workbook_updates": [{"dataset": dataset, "version": 0}]}
    assert client.post("/api/export", json=invalid_inventory_batch).status_code == 400
    inventory_batch_agent = client.post("/api/agent/call", headers={"Authorization": f"Bearer {client.services.token}"},
        json={"name": "data_export", "arguments": invalid_inventory_batch})
    assert inventory_batch_agent.status_code == 400
    response = client.post("/api/export", json={"dataset": dataset, "format": "xlsx", "mode": "preserve_workbook", "version": 1})
    assert response.status_code == 200, response.text
    api_result = response.json()
    assert api_result["dataset_version"] == 1
    assert json.loads(Path(api_result["receipt_path"]).read_text(encoding="utf-8"))["analysis_log_id"] == api_result["log_id"]

    client.services.undo(dataset, steps=1)
    agent = client.post("/api/agent/call", headers={"Authorization": f"Bearer {client.services.token}"},
                        json={"name": "data_export", "arguments": {"dataset": dataset, "format": "xlsx",
                              "mode": "preserve_workbook", "version": 0, "path": "agent-v0.xlsx"}})
    assert agent.status_code == 200, agent.text
    agent_result = agent.json()
    assert agent_result["mode"] == "preserve_workbook" and agent_result["dataset_version"] == 0
    original_version = openpyxl.load_workbook(agent_result["path"], data_only=False)
    assert original_version["Ventas"]["C2"].value == "café"
    original_version.close()
    assert client.services.log_get(agent_result["log_id"])["ok"] == 1
    flat = client.services.export(dataset, "xlsx", mode="flat")
    flat_wb = openpyxl.load_workbook(flat["path"], data_only=True)
    assert flat_wb.sheetnames == ["data"]
    flat_wb.close()


def test_multisheet_mcp_export_preflights_then_publishes_one_workbook(client, tmp_path, monkeypatch):
    source = _rich_workbook(tmp_path / "multi.xlsx", include_inventory=True)
    source_before = source.read_bytes()
    ingested = client.services.ingest_file(str(source), "multi", {"sheets": ["Ventas", "Inventario"]})
    sales, stock = [item["name"] for item in ingested["datasets"]]
    client.services.transform_apply(sales, "replace", {"column": "categoria", "pattern": "café", "replacement": "té"})
    client.services.transform_apply(stock, "replace", {"column": "categoria", "pattern": "inicial", "replacement": "revisado"})
    body = {"dataset": sales, "format": "xlsx", "mode": "preserve_workbook", "version": 1,
            "path": "multi-batch.xlsx", "workbook_updates": [{"dataset": stock, "version": 1}]}
    batch_started = time.perf_counter()
    response = client.post("/api/agent/call", headers={"Authorization": f"Bearer {client.services.token}"},
                           json={"name": "data_export", "arguments": body})
    batch_elapsed = time.perf_counter() - batch_started
    assert response.status_code == 200, response.text
    result = response.json()
    output = Path(result["path"])
    receipt_path = Path(result["receipt_path"])
    receipt_bytes = receipt_path.read_bytes()
    receipt = json.loads(receipt_bytes)
    assert [item["sheet"] for item in receipt["sheets"]] == ["Ventas", "Inventario"]
    assert [item["dataset_version"] for item in receipt["sheets"]] == [1, 1]
    assert receipt["cells_written"] == sum(item["cells_written"] for item in receipt["sheets"]) == 3
    log_entry = client.services.log_get(result["log_id"])
    assert "Ventas v1" in log_entry["output_summary"] and "Inventario v1" in log_entry["output_summary"]
    assert set(receipt["changed_package_parts"]) == {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml", "xl/workbook.xml"}
    wb = openpyxl.load_workbook(output, data_only=False)
    assert wb.sheetnames == ["Ventas", "Inventario", "Resumen", "Notas"]
    assert wb["Ventas"]["C2"].value == "té" and wb["Inventario"]["C2"].value == "revisado"
    assert wb["Ventas"]["E2"].data_type == "s" and wb["Ventas"]["F2"].value == "00123"
    assert wb["Resumen"]["B1"].value == "=SUM(Ventas!B2:B4)"
    assert wb["Notas"]["A2"].comment.text == "Comentario no afectado."
    assert len(wb["Notas"]._charts) == 1
    assert wb["Inventario"].tables["InventoryTable"].ref == "A1:C4"
    wb.close()
    assert source.read_bytes() == source_before
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(output) as dst:
        for name in src.namelist():
            if name not in receipt["changed_package_parts"]:
                assert src.read(name) == dst.read(name), name

    # REST uses the same multi-sheet contract; repeat exports are deterministic.
    rest_body = {**body, "path": "multi-rest.xlsx"}
    rest = client.post("/api/export", json=rest_body)
    assert rest.status_code == 200, rest.text
    assert rest.json()["output_sha256"] == result["output_sha256"]

    # Record measured batch-versus-single effects without asserting a speedup.
    singles_started = time.perf_counter()
    single_sales_response = client.post("/api/agent/call", headers={"Authorization": f"Bearer {client.services.token}"},
        json={"name": "data_export", "arguments": {"dataset": sales, "format": "xlsx", "mode": "preserve_workbook", "version": 1, "path": "single-sales.xlsx"}})
    single_stock_response = client.post("/api/agent/call", headers={"Authorization": f"Bearer {client.services.token}"},
        json={"name": "data_export", "arguments": {"dataset": stock, "format": "xlsx", "mode": "preserve_workbook", "version": 1, "path": "single-stock.xlsx"}})
    singles_elapsed = time.perf_counter() - singles_started
    assert single_sales_response.status_code == single_stock_response.status_code == 200
    single_sales, single_stock = single_sales_response.json(), single_stock_response.json()
    assert len(receipt["sheets"]) == 2 and single_sales["dataset"] == sales and single_stock["dataset"] == stock
    with zipfile.ZipFile(source) as package:
        package_bytes = sum(info.file_size for info in package.infolist())
    batch_artifacts = [output, receipt_path]
    single_artifacts = [Path(single_sales["path"]), Path(single_sales["receipt_path"]),
                        Path(single_stock["path"]), Path(single_stock["receipt_path"])]
    evidence_path = os.environ.get("NIGHTINGALE_MULTISHEET_EVIDENCE")
    if evidence_path:
        evidence = {"source_bytes": source.stat().st_size, "uncompressed_package_bytes": package_bytes,
                    "batch": {"mcp_calls": 1, "elapsed_seconds": batch_elapsed,
                              "workbooks": 1, "artifacts": len(batch_artifacts),
                              "bytes_written": sum(p.stat().st_size for p in batch_artifacts)},
                    "individual": {"mcp_calls": 2, "elapsed_seconds": singles_elapsed,
                                   "workbooks": 2, "artifacts": len(single_artifacts),
                                   "bytes_written": sum(p.stat().st_size for p in single_artifacts)},
                    "interpretation": "One real fixture measurement only; no performance speedup claim."}
        Path(evidence_path).parent.mkdir(parents=True, exist_ok=True)
        Path(evidence_path).write_text(json.dumps(evidence, indent=2), encoding="utf-8")

    # An invalid second version must not replace either an existing workbook or its receipt.
    prior_output, prior_receipt = output.read_bytes(), receipt_path.read_bytes()
    bad = {**body, "workbook_updates": [{"dataset": stock, "version": 999}]}
    rejected = client.post("/api/export", json=bad)
    assert rejected.status_code >= 400
    assert output.read_bytes() == prior_output and receipt_path.read_bytes() == prior_receipt

    # An invalid data shape in a later sheet is also found before publication.
    client.services.transform_apply(stock, "filter", {"expr": "sku = 'A-01'"})
    bad_shape = {**body, "workbook_updates": [{"dataset": stock, "version": 2}]}
    rejected_shape = client.post("/api/export", json=bad_shape)
    assert rejected_shape.status_code >= 400
    assert output.read_bytes() == prior_output and receipt_path.read_bytes() == prior_receipt

    duplicate = {**body, "workbook_updates": [{"dataset": stock, "version": 1},
                                                  {"dataset": stock, "version": 1}]}
    rejected_duplicate = client.post("/api/export", json=duplicate)
    assert rejected_duplicate.status_code >= 400
    assert output.read_bytes() == prior_output and receipt_path.read_bytes() == prior_receipt

    # A malformed second patched worksheet is rejected before publication too.
    from nightingale import workbook_export
    real_patch = workbook_export._patch_worksheet_cells
    patch_count = 0
    def malformed_second_sheet(xml, replacements, date_system):
        nonlocal patch_count
        patch_count += 1
        if patch_count == 2:
            return "<worksheet"
        return real_patch(xml, replacements, date_system)
    monkeypatch.setattr(workbook_export, "_patch_worksheet_cells", malformed_second_sheet)
    malformed = client.post("/api/export", json={**body, "path": "multi-batch.xlsx"})
    assert malformed.status_code == 400
    assert output.read_bytes() == prior_output and receipt_path.read_bytes() == prior_receipt

    catalog = next(item for item in tool_catalog() if item["name"] == "data_export")
    assert "workbook_updates" in catalog["inputSchema"]["properties"]


def test_preserving_export_uses_source_1904_date_epoch(services, tmp_path):
    source = _rich_workbook(tmp_path / "epoch-1904.xlsx", date1904=True)
    source_inventory = inventory(source)
    assert source_inventory["date_system"] == "1904"
    dataset = _ingest_target(services, source)
    result = services.export(dataset, "xlsx", mode="preserve_workbook")
    output = Path(result["path"])
    receipt = json.loads(Path(result["receipt_path"]).read_text(encoding="utf-8"))
    assert receipt["date_system"] == "1904"
    exported = openpyxl.load_workbook(output, data_only=False)
    assert exported.epoch == openpyxl.utils.datetime.CALENDAR_MAC_1904
    assert exported["Ventas"]["A2"].value == openpyxl.utils.datetime.from_ISO8601("2026-01-03T00:00:00")
    assert exported["Ventas"]["A2"].data_type == "d"
    exported.close()
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        match = re.search(r'<c\b[^>]*\br="A2"[^>]*>.*?<v>([^<]+)</v>', xml)
        assert match is not None
        expected_serial = openpyxl.utils.datetime.to_excel(
            openpyxl.utils.datetime.from_ISO8601("2026-01-03T00:00:00"),
            openpyxl.utils.datetime.CALENDAR_MAC_1904,
        )
        assert float(match.group(1)) == expected_serial


def test_1900_epoch_matches_excel_serials_before_and_after_leap_bug():
    assert float(_excel_value(openpyxl.utils.datetime.from_ISO8601("1900-01-01T00:00:00"), "1900")[0]) == 1
    assert float(_excel_value(openpyxl.utils.datetime.from_ISO8601("1900-02-28T00:00:00"), "1900")[0]) == 59
    assert float(_excel_value(openpyxl.utils.datetime.from_ISO8601("1900-03-01T00:00:00"), "1900")[0]) == 61


@pytest.mark.parametrize("omit_calc_pr", [False, True], ids=["existing-calcPr", "missing-calcPr"])
def test_prefixed_workbook_and_worksheet_xml_preserve_source_and_reopen(services, tmp_path, omit_calc_pr):
    source = _rich_workbook(tmp_path / f"prefixed-{omit_calc_pr}.xlsx")
    _prefix_workbook_xml(source, omit_calc_pr=omit_calc_pr)
    original = source.read_bytes()
    with zipfile.ZipFile(source) as archive:
        before_workbook = archive.read("xl/workbook.xml").decode("utf-8")
        assert "<x:workbook" in before_workbook
        assert ("<x:calcPr" not in before_workbook) is omit_calc_pr
    dataset = _ingest_target(services, source)
    services.transform_apply(dataset, "derive", {"name": "importe", "expr": "importe + 10"})
    services.transform_apply(dataset, "derive", {"name": "activo", "expr": "NOT activo"})
    services.transform_apply(dataset, "replace", {"column": "categoria", "pattern": "café", "replacement": "té"})
    result = services.export(dataset, "xlsx", mode="preserve_workbook")
    with zipfile.ZipFile(result["path"]) as archive:
        workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
        worksheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert "<x:calcPr" in workbook_xml
        assert 'fullCalcOnLoad="1"' in workbook_xml and 'forceFullCalc="1"' in workbook_xml
        assert "<x:worksheet" in worksheet_xml and "<x:sheetData" in worksheet_xml
        assert "<x:c r=\"C2\"" in worksheet_xml
        assert "<x:is><x:t>té</x:t></x:is>" in worksheet_xml
        parsed_sheet = openpyxl.xml.functions.fromstring(worksheet_xml)
        main_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        assert parsed_sheet.tag == main_ns + "worksheet"
        assert parsed_sheet.find(f".//{main_ns}c[@r='B2']/{main_ns}v").text == "22.5"
        assert parsed_sheet.find(f".//{main_ns}c[@r='D2']/{main_ns}v").text == "0"
    reopened = openpyxl.load_workbook(result["path"], data_only=False)
    assert reopened["Ventas"]["C2"].value == "té"
    assert reopened["Ventas"]["B2"].value == 22.5
    assert reopened["Ventas"]["D2"].value is False
    reopened.close()
    assert source.read_bytes() == original


def test_one_pass_sheet_patch_handles_sparse_self_closing_cells_and_1400_rows():
    sparse = ("<worksheet xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'>"
              "<sheetData><row r='1'><c r='A1' t='inlineStr'><is><t>head</t></is></c></row>"
              "<row r='2'><c r='A2'><v>1</v></c></row>"
              "<row r='3'><c r='A3'><v>2</v></c><c r='C3' s='7'/></row></sheetData></worksheet>")
    patched = _patch_worksheet_cells(sparse, {"C3": "now filled"}, "1900")
    parsed = openpyxl.xml.functions.fromstring(patched)
    cell = parsed.find(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c[@r='C3']")
    assert cell is not None and cell.attrib["s"] == "7" and cell.attrib["t"] == "inlineStr"
    assert cell.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}is/{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t").text == "now filled"

    row_count = 1400
    pieces = ["<worksheet xmlns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'><sheetData>"]
    replacements = {}
    for row in range(1, row_count + 1):
        pieces.append(f"<row r='{row}'>" + "".join(f"<c r='{col}{row}'><v>{row}</v></c>" for col in "ABCDEFGH") + "</row>")
        replacements[f"G{row}"] = f"updated-{row}"
    pieces.append("</sheetData></worksheet>")
    started = time.perf_counter()
    large = _patch_worksheet_cells("".join(pieces), replacements, "1900")
    elapsed = time.perf_counter() - started
    large_root = openpyxl.xml.functions.fromstring(large)
    assert len(large_root.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c[@t='inlineStr']")) == row_count
    assert elapsed < 10.0, f"one-pass 1,400-row worksheet patch took {elapsed:.2f}s"


def test_changed_text_with_carriage_returns_reopens_and_illegal_values_refuse(services, tmp_path):
    source = _rich_workbook(tmp_path / "text-values.xlsx")
    source_bytes = source.read_bytes()
    dataset = _ingest_target(services, source)
    replacement = "  café\r\nline  "
    services.transform_apply(dataset, "replace", {"column": "categoria", "pattern": "café", "replacement": replacement})
    result = services.export(dataset, "xlsx", mode="preserve_workbook")
    output_path = Path(result["path"])
    reopened = openpyxl.load_workbook(output_path, data_only=False)
    assert reopened["Ventas"]["C2"].value == replacement
    reopened.close()
    with zipfile.ZipFile(output_path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"&#13;" in xml and b'xml:space="preserve"' in xml
    assert source.read_bytes() == source_bytes

    with pytest.raises(Exception, match="non-finite floating-point"):
        _cell_xml('<c r="A1">', float("nan"), "1900")
    with pytest.raises(Exception, match="non-finite Decimal"):
        from decimal import Decimal
        _cell_xml('<c r="A1">', Decimal("Infinity"), "1900")
    with pytest.raises(Exception, match="XML-illegal control"):
        _patch_worksheet_cells("<worksheet><sheetData><row r='1'></row></sheetData></worksheet>",
                               {"A1": "bad\x01text"}, "1900")


def test_formula_in_target_range_is_refused_and_failures_are_logged(services, tmp_path):
    source = _rich_workbook(tmp_path / "formula.xlsx", target_formula=True)
    dataset = _ingest_target(services, source)
    services.transform_apply(dataset, "replace", {"column": "categoria", "pattern": "café", "replacement": "té"})
    kept = services.export(dataset, "xlsx", mode="preserve_workbook")
    keep_wb = openpyxl.load_workbook(kept["path"], data_only=False)
    assert keep_wb["Ventas"]["B2"].value == "=6+6.5"
    keep_wb.close()

    version = services._version_row(services._dataset_row(dataset))
    with services.engine.lock() as conn:
        conn.execute(f'UPDATE "{version["table_name"]}" SET "importe" = 99 WHERE rowid = 0')
    with pytest.raises(Exception, match="changed value would overwrite formula cell B2"):
        services.export(dataset, "xlsx", mode="preserve_workbook")
    assert any(entry["ok"] == 0 and "would overwrite formula cell B2" in (entry["error"] or "")
               for entry in services.log_search(dataset=dataset)["entries"])


def test_validation_comments_and_hyperlinks_survive_same_range_edit(services, tmp_path):
    source = _rich_workbook(tmp_path / "validated.xlsx", target_validation=True)
    source_inventory = inventory(source)
    assert source_inventory["sheets"][0]["comments"][0]["refs"] == ["C2"]
    dataset = _ingest_target(services, source)
    services.transform_apply(dataset, "replace", {"column": "categoria", "pattern": "café", "replacement": "té"})
    result = services.export(dataset, "xlsx", mode="preserve_workbook")
    output = openpyxl.load_workbook(result["path"], data_only=False)
    sheet = output["Ventas"]
    assert sheet["C2"].value == "té"
    assert sheet["C2"].comment.text == "Keep this note while editing."
    assert sheet["C2"].hyperlink.target == "https://example.invalid/category"
    assert str(sheet.data_validations.dataValidation[0].sqref) == "C2:C4"
    assert result["preserved_metadata"]["comment_count"] == 1
    assert result["preserved_metadata"]["hyperlink_count"] == 1
    assert result["preserved_metadata"]["data_validation_count"] == 1
    assert "not evaluated" in result["preserved_metadata"]["data_validation_policy"]
    output.close()


def test_source_snapshot_change_and_unsafe_shape_refuse_without_output(services, tmp_path):
    source = _rich_workbook(tmp_path / "changed.xlsx")
    dataset = _ingest_target(services, source)
    original = source.read_bytes()
    source.write_bytes(original + b" ")
    with pytest.raises(Exception, match="changed after ingestion"):
        services.export(dataset, "xlsx", mode="preserve_workbook")
    assert not list(services.config.exports_dir.glob("*workbook.xlsx"))

    source.write_bytes(original)
    # A row-count-changing version cannot be fitted back into a fixed workbook range.
    services.transform_apply(dataset, "filter", {"expr": 'categoria <> \'niño\''})
    with pytest.raises(Exception, match="requires the imported row count"):
        services.export(dataset, "xlsx", mode="preserve_workbook")
    assert not list(services.config.exports_dir.glob("*workbook.xlsx"))
